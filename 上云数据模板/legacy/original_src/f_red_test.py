#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""验证 超差 -> Result=F 且单元格变红 的逻辑是否真实生效。"""
import os
import tempfile

import converter
from openpyxl import load_workbook

data_rows = [
    # 1) 超上限 -> 期望 F + 红
    {"src_row": 0, "item_no": "1.0", "spec": "10", "max": "10.2", "min": "9.8",
     "instr": "YK", "insp": "尺寸", "readings": ["11.5"]},
    # 2) 合格 -> 期望 P
    {"src_row": 1, "item_no": "2.0", "spec": "5", "max": "5.2", "min": "4.8",
     "instr": "YK", "insp": "尺寸", "readings": ["5.0"]},
    # 3) 低于下限 -> 期望 F + 红
    {"src_row": 2, "item_no": "3.0", "spec": "8", "max": "8.2", "min": "7.8",
     "instr": "YK", "insp": "尺寸", "readings": ["7.0"]},
    # 4) 全 Pass（技术要求） -> 期望 P
    {"src_row": 3, "item_no": "A", "spec": "未注公差", "max": "/", "min": "/",
     "instr": "目测", "insp": "技术要求", "readings": ["Pass"]},
]

hinfo = {
    "title": "Inspection Report", "dwg_no": "TEST001", "part_name": "测试件",
    "customer_code": "X", "po": "", "rev": "A", "temp": "20℃",
    "inspected_by": "检验:张三", "reviewed_by": "审核:李四",
    "inspected_date": "2026/8/13", "reviewed_date": "2026/8/13",
    "total_qty": 1, "sample_qty": 1,
}

d = tempfile.mkdtemp()
inp = os.path.join(d, "TEST001_测试件_1件_26.8.13.xlsx")
res = converter._build_output_xlsx(data_rows, {}, inp, d, None,
                                   sv=None, header_info=hinfo, total_qty=1)
out_path = res[0] if isinstance(res, tuple) else res["out_path"]

wb = load_workbook(out_path)
ws = wb.active
# total_qty=1 -> Result 列 = L(12)
print("OUTPUT:", out_path)
ok = True
for r in range(11, ws.max_row + 1):
    v = ws.cell(row=r, column=12).value
    cell = ws.cell(row=r, column=12)
    fill = cell.fill
    is_red = (fill and fill.fill_type == "solid" and fill.fgColor
              and "FF0000" in str(fill.fgColor.rgb))
    font_red = cell.font and cell.font.color and "FFFFFF" in str(cell.font.color.rgb)
    item = ws.cell(row=r, column=1).value
    tag = ""
    if v == "F":
        tag = "RED" if (is_red and font_red) else "!!NO-RED"
        if not (is_red and font_red):
            ok = False
    print(f"  item={item!r:>8} Result={v!r:>4} fill={str(fill.fgColor.rgb) if fill and fill.fgColor else '-':>10} {tag}")

print("RESULT:", "PASS 超差行均 F+红" if ok else "FAIL 红色/字体未生效")
