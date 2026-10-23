# -*- coding: utf-8 -*-
"""专门验证 .xls 支持（Excel COM 转临时 .xlsx 后再走原流程）。"""
import os, sys, time, shutil

HERE = os.path.dirname(os.path.abspath(__file__))
TMP = os.path.join(HERE, f"sandbox_xls_{int(time.time()*1000)}")
sys.path.insert(0, os.path.dirname(HERE))

import openpyxl
from extract_core.api import run_batch
from extract_core.filename_parser import parse_report_filename
from extract_core.errors import ExtractError
from extract_core.report_reader import detect_layout

BH = "930XLS-DM"
DESC = "xls测试件"
DATE = "26.4.09"


def build_xlsx_template(out_dir, used, total, nsheets=3, name=None):
    """造一份与真实测量报告同构的 .xlsx 模板。"""
    fname = name or f"{BH}_{DESC}_{total}件_{DATE}-{used}(1).xlsx"
    path = os.path.join(out_dir, fname)
    wb = openpyxl.Workbook()
    for k in range(1, nsheets + 1):
        ws = wb.create_sheet(f"2026.4.9（{total}件）-{k}")
        ws["A1"] = "Inspection Report\n检验报告"
        ws["A3"] = "DWG No:\n图号"; ws["B3"] = BH
        ws["D3"] = "Part Description:\n零件名称"; ws["F3"] = DESC
        ws["A9"] = "Item No\n（编号）"; ws["B9"] = "Specified Parameter"
        ws["E9"] = "Instrument used（检具）"
        ws["F9"] = "Total Qty\n（总数）:"; ws["H9"] = total
        ws["I9"] = "Sample Qty\n（检查数）:"; ws["L9"] = 5
        ws["F10"] = "Actual Readings（测量数据）"; ws["L10"] = "Result（结果）"
        ws["B11"] = "公称尺寸\nDimension"; ws["C11"] = "上限值\nMax"; ws["D11"] = "下限值\nMin"
        seqs = [5 * (k - 1) + i for i in range(1, 6)]
        for off, g in enumerate(seqs):
            ws.cell(row=11, column=6 + off, value=g)
        ws["K11"] = "(P合格 /F不合格)"; ws["L11"] = "最大值"; ws["M11"] = "最小值"
        for ridx in range(10):
            r = 12 + ridx
            ws.cell(row=r, column=1, value=ridx + 1)
            ws.cell(row=r, column=2, value=f"spec{r}")
            ws.cell(row=r, column=3, value=10.0); ws.cell(row=r, column=4, value=9.0)
            ws.cell(row=r, column=5, value="测高仪")
            for off, g in enumerate(seqs):
                ws.cell(row=r, column=6 + off, value=float(f"{g}.{r}"))
            ws.cell(row=r, column=11, value="P")
            ws.cell(row=r, column=12, value=0)
            ws.cell(row=r, column=13, value=0)
    wb.remove(wb["Sheet"])
    wb.save(path)
    return path


def xlsx_to_xls(xlsx_path):
    """用 Excel COM 把 .xlsx 转成 .xls（BIFF8）。"""
    import pythoncom, win32com.client
    xls_path = os.path.splitext(xlsx_path)[0] + ".xls"
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False; xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(os.path.abspath(xlsx_path), ReadOnly=True, UpdateLinks=0)
        wb.SaveAs(os.path.abspath(xls_path), FileFormat=56)  # 56 = xlExcel8 (.xls)
        wb.Close(False)
    finally:
        xl.Quit()
        pythoncom.CoUninitialize()
    return xls_path


def build_bom(out_dir, rows):
    path = os.path.join(out_dir, f"BOM_{int(time.time()*1000)}.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "子阶物料"; ws["B1"] = "顶层物料"; ws["C1"] = "图号"
    ws["D1"] = "组件描述"; ws["E1"] = "组件数量"; ws["F1"] = "图纸+尺寸报告+材质证明"
    for i, (bh, qty) in enumerate(rows):
        r = 2 + i; ws[f"C{r}"] = bh; ws[f"E{r}"] = qty; ws[f"F{r}"] = 1
    wb.save(path)
    return path


def verify_output(out_path, K, used, total):
    wb = openpyxl.load_workbook(out_path, data_only=True)
    assert wb.sheetnames[0] == DATE, f"表名应为 {DATE}，实为 {wb.sheetnames[0]}"
    ws = wb[wb.sheetnames[0]]
    seq = [ws.cell(row=11, column=6 + j).value for j in range(5)]
    assert seq[:K] == list(range(used + 1, used + K + 1)), f"序号行前{K}应为{used+1}..{used+K}，实为{seq[:K]}"
    assert all(v is None for v in seq[K:]), "序号行多余应留空"
    assert ws.cell(row=9, column=8).value == K and ws.cell(row=9, column=12).value == K
    print(f"  [校验通过] {os.path.basename(out_path)}：序号={seq[:K]}, 总数/检查={K}")


def cur_suffix(path):
    import re
    base = os.path.basename(path)
    base = re.sub(r"\(\d+\)", "", base)
    m = re.search(r"-(\d+)\.(xlsx|xls)$", base)
    return int(m.group(1)) if m else None


def main():
    if os.path.exists(TMP):
        shutil.rmtree(TMP)
    os.makedirs(TMP)
    fails = []

    # ===== 用例A：detect_layout 在 .xls 上不再崩溃 =====
    print("=== A. .xls 布局探测（COM 转临时 xlsx）===")
    xlsA = xlsx_to_xls(build_xlsx_template(TMP, used=2, total=10))
    assert os.path.exists(xlsA), "未生成 .xls"
    try:
        info = detect_layout(xlsA)
        assert info.sheets and info.sample_map, "布局探测应找到工作表与测量列"
        print(f"  [OK] .xls 探测成功：{len(info.sheets)} 表, {len(info.sample_map)} 件")
    except Exception as e:
        fails.append(f"A detect_layout .xls 失败: {e}")
        print("  [FAIL]", e)

    # ===== 用例B：端到端从 .xls 抽取（BOM 2件）=====
    print("=== B. 端到端：.xls 作为来源抽取 2 件 ===")
    xlsB = xlsx_to_xls(build_xlsx_template(TMP, used=2, total=10))
    bom = build_bom(TMP, [(BH, 2)])
    s = run_batch(bom, [xlsB], out_dir=TMP)
    print("  ok/fail:", s["ok_count"], s["fail_count"])
    if s["ok_count"] == 1 and s["fail_count"] == 0:
        t = s["tasks"][0]
        verify_output(t["output"], 2, 2, 10)
        # 来源 .xls 应被改名推进到 -4
        after = t["source_after"]
        assert after.lower().endswith(".xls"), f"来源应为 .xls，实为 {after}"
        assert cur_suffix(after) == 4, f"来源应推进到 -4，实为 -{cur_suffix(after)}"
        print(f"  [OK] 生成 {os.path.basename(t['output'])}（.xlsx），来源改名 {os.path.basename(after)}")
    else:
        fails.append(f"B 端到端 .xls 失败: {s}")
        print("  [FAIL]", s)

    # ===== 用例C：.xls 图号未匹配 BOM -> NO_BOM_MATCH，整批不崩溃 =====
    print("=== C. .xls 图号未匹配 -> NO_BOM_MATCH（隔离，不崩溃）===")
    xlsC = xlsx_to_xls(build_xlsx_template(TMP, used=2, total=10, name="930OTHER-DM_xls其它_10件_26.4.09-2(1).xlsx"))
    # 同时放一个能匹配的 .xls，验证前者失败不影响后者成功
    xlsC2 = xlsx_to_xls(build_xlsx_template(TMP, used=1, total=10))
    other = [(f"NO-{i:03d}", 2) for i in range(20)]  # 无关图号，无 930OTHER-DM 也无 930XLS-DM
    bomC = build_bom(TMP, other + [(BH, 2)])  # 仅含 930XLS-DM
    try:
        s = run_batch(bomC, [xlsC, xlsC2], out_dir=TMP)
        codes = [(t.get("kind"), t.get("error_code")) for t in s["tasks"]]
        print("  tasks:", codes)
        # xlsC 应 NO_BOM_MATCH(file)，xlsC2 应 ok(task)
        file_err = [t for t in s["tasks"] if t.get("kind") == "file"]
        task_ok = [t for t in s["tasks"] if t.get("ok")]
        assert any(t.get("error_code") == "NO_BOM_MATCH" for t in file_err), "缺 NO_BOM_MATCH"
        assert len(task_ok) == 1, f"应能成功1个，实为 {len(task_ok)}"
        print("  [OK] .xls 未匹配被隔离为 NO_BOM_MATCH，同批另一 .xls 正常完成")
    except Exception as e:
        fails.append(f"C 隔离测试崩溃: {type(e).__name__}: {e}")
        print("  [FAIL] 整批崩溃:", e)

    # ===== 用例D：非常规命名仍报错（优化文案）=====
    print("=== D. 非常规命名 91010070...XLSX -> FILE_PARSE_ERROR（新文案）===")
    try:
        parse_report_filename(os.path.join(TMP, "91010070_E.1_20260814_Pd主腔体出货报告清单1(1)(3).XLSX"))
        fails.append("D 应报错却通过")
        print("  [FAIL] 未报错")
    except ExtractError as e:
        print(f"  [OK] {e.code}: {e.reason}")

    # ===== 用例E：.doc 仍报错 =====
    print("=== E. 9.10.doc -> FILE_PARSE_ERROR ===")
    try:
        parse_report_filename(os.path.join(TMP, "9.10.doc"))
        fails.append("E 应报错却通过")
        print("  [FAIL] 未报错")
    except ExtractError as e:
        print(f"  [OK] {e.code}: {e.reason}")

    print("\n==== 结果 ====")
    if fails:
        print("❌ 失败项:")
        for f in fails:
            print("  -", f)
        raise SystemExit(1)
    print("✅ .xls 支持全部验证通过（探测/抽取/隔离 + 非常规命名与 .doc 报错）")
    print("  沙箱:", TMP)


if __name__ == "__main__":
    main()
