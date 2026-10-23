# -*- coding: utf-8 -*-
"""数字键匹配专项测试（v5 新增的第 4 级匹配）。

覆盖：
  A. 报告只写图号的一部分（BOM 图号带 JH-YCL- 前后缀）-> 命中
  B. 同数字多行 -> 用组件描述消歧（选中描述与文件名一致的那行）
  C. 同数字多行且描述都无法消歧 -> BOM_MULTI_MATCH（不瞎选）
  D. 数字键太短（<5 位，标准件目录号）-> 不参与，仍 NO_BOM_MATCH
  E. 回归：原有 ①②③ 三级行为完全不变
"""
import os
import sys
import time

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TMP = os.path.join(HERE, f"numkey_{int(time.time() * 1000)}")
sys.path.insert(0, os.path.dirname(HERE))
os.makedirs(TMP, exist_ok=True)

from extract_core.bom_reader import read_bom_ex, match_front
from extract_core.errors import ExtractError

fails = []


def build_bom(name, rows):
    """PM-DM 型 BOM：C=图号 D=组件描述 E=组件数量。rows=(bh, desc, qty)"""
    path = os.path.join(TMP, name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "子阶物料"; ws["B1"] = "顶层物料"; ws["C1"] = "图号"
    ws["D1"] = "组件描述"; ws["E1"] = "组件数量"; ws["F1"] = "图纸"
    for i, (bh, desc, qty) in enumerate(rows):
        r = 2 + i
        ws[f"C{r}"] = bh
        ws[f"D{r}"] = desc
        ws[f"E{r}"] = qty
    wb.save(path)
    return path


def check(case, front, bom_rows, want_qty=None, want_desc=None, want_code=None,
          want_empty=False):
    bom = build_bom(f"BOM_{case}.xlsx", bom_rows)
    rows, _meta = read_bom_ex(bom)
    try:
        hits = match_front(front, rows)
        got_code = None
    except ExtractError as e:
        hits, got_code = [], e.code
    if want_empty:
        if got_code is None and not hits:
            print(f"  [OK] {case}: front={front!r} -> 未命中（符合预期）")
        else:
            fails.append(f"{case}: 期望未命中，实为 code={got_code} hits={hits}")
            print(f"  [FAIL] {case}: 期望未命中，实为 code={got_code} hits={hits}")
        return
    if want_code:
        if got_code == want_code:
            print(f"  [OK] {case}: front={front!r} -> {got_code}")
        else:
            fails.append(f"{case}: 期望 {want_code}，实为 {got_code}（hits={hits}）")
            print(f"  [FAIL] {case}: 期望 {want_code}，实为 {got_code}")
        return
    if got_code:
        fails.append(f"{case}: 意外报错 {got_code}")
        print(f"  [FAIL] {case}: 意外报错 {got_code}")
        return
    if len(hits) == 1 and hits[0].qty == want_qty and (
            want_desc is None or hits[0].desc == want_desc):
        print(f"  [OK] {case}: front={front!r} -> 图号={hits[0].bh!r} 描述={hits[0].desc!r} 件数={hits[0].qty}")
    else:
        fails.append(f"{case}: 期望单行 件数={want_qty} 描述={want_desc}，实为 {hits}")
        print(f"  [FAIL] {case}: 实为 {hits}")


print("=== A. 报告只写图号的一部分（BOM 图号带 JH-YCL- 前缀）===")
check("A", "93010410_保护盖1",
      [("JH-YCL-93010410（按图加工）", "保护盖1", 4)],
      want_qty=4, want_desc="保护盖1")

print("=== B. 同数字多行 -> 描述消歧 ===")
check("B", "93019282_翻转横撑",
      [("93019282-DM", "弹簧压板", 3),
       ("JH-YCL-93019282-DM（按图加工）", "翻转横撑", 1)],
      want_qty=1, want_desc="翻转横撑")

print("=== C. 同数字多行且无法消歧 -> BOM_MULTI_MATCH ===")
check("C", "90000001_别的零件",
      [("90000001-A", "AAA", 2), ("90000001-B", "BBB", 3)],
      want_code="BOM_MULTI_MATCH")

print("=== D. 数字键太短（<5 位）-> 不参与 ===")
check("D", "1191_圆柱销",
      [("GB/T 119.1", "圆柱销", 5)],
      want_empty=True)

print("=== E. 回归：原有三级行为不变 ===")
check("E1", "93011182_气路法兰", [("93011182", "气路法兰", 1)], want_qty=1)          # ① 拼接全等
check("E2", "93011182_别的描述", [("93011182", "气路法兰", 1)], want_qty=1)          # ② 图号前缀
check("E3", "93126082_远端轴套", [("", "远端轴套_93126082_Pla,PTFE-V0", 2)], want_qty=2)  # ③ 分段集合

print("\n==== 结果 ====")
if fails:
    print("❌ 失败项:")
    for f in fails:
        print("  -", f)
    raise SystemExit(1)
print("✅ 数字键匹配全部通过（部分图号 / 描述消歧 / 多义保护 / 短键不参与 / 三级回归）")
print("  沙箱:", TMP)
