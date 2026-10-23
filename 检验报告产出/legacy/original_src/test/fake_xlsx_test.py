# -*- coding: utf-8 -*-
"""假 .xlsx 自动修复回归测试。

场景：WPS 按老 .xls(BIFF) 写盘却套 .xlsx 扩展名（Excel 拒绝打开）。
验证：_is_fake_xlsx 识别、extract_identity 自动转换读取、detect_layout 自动转换探测。
自包含：用 openpyxl 生成合成报告 -> COM 存成老 .xls -> 复制成 .xlsx 名（假文件）。
依赖本机 Excel COM（同 xls_support_test）。
"""
import os
import shutil
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from extract_core.report_reader import (
    _is_fake_xlsx,
    detect_layout,
    extract_identity,
)

BH = "930FAKE-DM"
DESC = "假格式测试件"
TOTAL = 5
DATE = "26.5.6"


def build_real_xlsx(path):
    """生成一张标准合成报告（真 .xlsx）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "检验报告"
    ws["A1"] = "Inspection Report\n检验报告"
    ws["A3"] = "DWG No:\n图号"; ws["B3"] = BH
    ws["D3"] = "Part Description:\n零件名称"; ws["F3"] = DESC
    ws["F9"] = "Total Qty\n（总数）:"; ws["H9"] = TOTAL
    ws["I9"] = "Sample Qty\n（检查数）:"; ws["K9"] = TOTAL
    ws["F10"] = "Actual Readings（测量数据）"; ws["K10"] = "Result（结果）"
    ws["B11"] = "公称尺寸\nDimension"; ws["C11"] = "上限值\nMax"; ws["D11"] = "下限值\nMin"
    for j in range(5):
        ws.cell(row=11, column=6 + j, value=j + 1)          # 序号 1..5（col6-10）
    ws.cell(row=11, column=11, value="(P合格 /F不合格)")     # 结果列表头（紧贴第5列）
    for ridx in range(10):
        r = 12 + ridx
        ws.cell(row=r, column=2, value=f"spec{ridx}")
        for j in range(5):
            ws.cell(row=r, column=6 + j, value=float(f"{j+1}.{ridx}"))
        ws.cell(row=r, column=11, value="P")
    wb.save(path)
    return path


def to_legacy_xls(src_xlsx, dst_xls):
    """用 Excel COM 把 .xlsx 另存为老 .xls(BIFF, FileFormat=56)。"""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(os.path.abspath(src_xlsx))
        try:
            wb.SaveAs(os.path.abspath(dst_xls), FileFormat=56)  # 56 = xlExcel8
        finally:
            wb.Close(SaveChanges=False)
    finally:
        xl.Quit()
        pythoncom.CoUninitialize()
    return dst_xls


def main():
    tmpdir = tempfile.mkdtemp(prefix="fake_xlsx_test_")
    try:
        real = os.path.join(tmpdir, f"{BH}_{DESC}_{TOTAL}件_{DATE}.xlsx")
        legacy = os.path.join(tmpdir, f"{BH}_{DESC}_{TOTAL}件_{DATE}.xls")
        fake = os.path.join(tmpdir, f"{BH}_{DESC}_{TOTAL}件_{DATE}(1).xlsx")  # 假：老内容+新扩展名
        build_real_xlsx(real)
        to_legacy_xls(real, legacy)
        shutil.copy2(legacy, fake)

        with open(fake, "rb") as f:
            head = f.read(8)
        assert head == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", f"假文件应为 OLE2 签名：{head.hex()}"

        # 1) 识别
        assert _is_fake_xlsx(fake) is True
        assert _is_fake_xlsx(real) is False

        # 2) extract_identity 自动转换读取
        dwg, part = extract_identity(fake)
        assert dwg == BH and part == DESC, (dwg, part)
        print("  [OK] extract_identity 自动转换 ->", (dwg, part))

        # 3) detect_layout 自动转换探测
        info = detect_layout(fake)
        assert len(info.sheets) == 1
        s = info.sheets[0]
        assert len(s.sample_cols) == 5 and s.result_col == 11, (s.sample_cols, s.result_col)
        assert list(info.sample_map) == [1, 2, 3, 4, 5]
        print("  [OK] detect_layout 自动转换 -> 5 列 / 结果列 11 / 件号 1..5")

        # 4) 真 .xlsx 不受影响
        info2 = detect_layout(real)
        assert len(info2.sheets) == 1 and list(info2.sample_map) == [1, 2, 3, 4, 5]
        print("  [OK] 真 .xlsx 路径不受影响")

        print("PASS fake_xlsx_test")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    main()
