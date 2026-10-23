# -*- coding: utf-8 -*-
"""e2e 验证（v6.2 唯一行口径）：BOM 两行【完全一样】-> 同图号多行 -> 文件级直接报错。

表一图号唯一；命中多行（含等价行）一律 BOM_MULTI_MATCH，不建任务、不回写、不动来源。
"""
import os, shutil, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from extract_core.api import run_batch

SANDBOX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sandbox_dup_e2e")
if os.path.isdir(SANDBOX):
    shutil.rmtree(SANDBOX)
os.makedirs(SANDBOX, exist_ok=True)

BH, DESC, DATE = "930EDGE-DM", "重复测试件", "26.4.09"

def build_report(out_dir, total=5, nsheets=5):
    fname = f"{BH}_{DESC}_{total}件_{DATE}.xlsx"
    path = os.path.join(out_dir, fname)
    wb = openpyxl.Workbook()
    for k in range(1, nsheets + 1):
        ws = wb.create_sheet(f"2026.4.9（{total}件）-{k}")
        ws["A1"] = "Inspection Report\n检验报告"
        ws["A3"] = "DWG No:\n图号"; ws["B3"] = BH
        ws["D3"] = "Part Description:\n零件名称"; ws["F3"] = DESC
        ws["A9"] = "Item No\n（编号）"; ws["B9"] = "Specified Parameter"
        ws["E9"] = "Instrument used（检具）"
        ws["F9"] = "Total Qty\n（总数）:"; ws["H9"] = total
        ws["I9"] = "Sample Qty\n（检查数）:"; ws["L9"] = 5
        ws["F10"] = "Actual Readings（测量数据）"; ws["L10"] = "Result（结果）"
        ws["B11"] = "公称尺寸\nDimension"; ws["C11"] = "上限值\nMax"; ws["D11"] = "下限值\nMin"
        seqs = [5 * (k - 1) + i for i in range(1, 6)]
        for off, g in enumerate(seqs):
            ws.cell(row=11, column=6 + off, value=g)
        ws["K11"] = "(P合格 /F不合格)"; ws["L11"] = "最大值"; ws["M11"] = "最小值"
        for ridx in range(10):
            r = 12 + ridx
            ws.cell(row=r, column=1, value=ridx + 1)
            ws.cell(row=r, column=2, value=f"spec{r}")
            ws.cell(row=r, column=3, value=10.0); ws.cell(row=r, column=4, value=9.0)
            ws.cell(row=r, column=5, value="测高仪")
            for off, g in enumerate(seqs):
                ws.cell(row=r, column=6 + off, value=float(f"{g}.{r}"))
            ws.cell(row=r, column=11, value="P")
            ws.cell(row=r, column=12, value=0)
            ws.cell(row=r, column=13, value=0)
    wb.remove(wb["Sheet"])
    wb.save(path)
    return path

# 1) 造 BOM：两行完全一样（图号+描述+数量都同）
bom_path = os.path.join(SANDBOX, "BOM_dup.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(["图号", "组件描述", "组件数量", "第13套"])
ws.append([BH, DESC, 2, None])   # 第2行
ws.append([BH, DESC, 2, None])   # 第3行（与第2行完全一致）
wb.save(bom_path)

report = build_report(SANDBOX)
out_dir = os.path.join(SANDBOX, "out")
res = run_batch(bom_path, [report], out_dir=out_dir, round_no=13)

print("== run_batch 结果 ==")
for t in res["tasks"]:
    print(f"  kind={t.get('kind')} ok={t.get('ok')} bh={t.get('bh')} "
          f"err={t.get('error_code')} reason={str(t.get('reason'))[:120]}")

fails = [t for t in res["tasks"] if not t.get("ok")]
assert len(fails) == 1, f"应恰有 1 个文件级失败，实际 {len(fails)}"
assert fails[0]["error_code"] == "BOM_MULTI_MATCH", f"应为 BOM_MULTI_MATCH，实际 {fails[0]['error_code']}"
assert all("重复测试件" in str(t.get("reason", "")) for t in fails), "报错应列出候选行（含描述）"

# 无成功任务
assert res["ok_count"] == 0, "不应有任何成功任务"

# BOM 未回写（第13套列两行都为空）
wb2 = openpyxl.load_workbook(bom_path, data_only=True)
ws2 = wb2["Sheet1"]
assert ws2.cell(row=2, column=4).value is None, "第2行不应回写"
assert ws2.cell(row=3, column=4).value is None, "第3行不应回写"
print("  BOM 未回写：第2/第3行「第13套」均保持为空")

# 来源文件未动（无 -used 推进）
import glob
afters = [os.path.basename(p) for p in glob.glob(os.path.join(SANDBOX, "930EDGE-DM*.xlsx"))]
print(f"  来源保持原样：{afters}")
assert any(a.endswith(f"{DATE}.xlsx") for a in afters), f"来源不应被改名：{afters}"

print("\n==== 两行完全一样 -> 同图号多行直接报错：端到端验证通过 ====")
