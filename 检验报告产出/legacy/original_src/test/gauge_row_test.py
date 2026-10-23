# -*- coding: utf-8 -*-
"""验证：数据区末尾以「检具代码」行为界，不再受 400 行硬截断。"""
import os, sys, shutil, traceback

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import Font
from extract_core.report_reader import detect_layout
from extract_core.filename_parser import parse_report_filename

TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_gauge_tmp")
os.makedirs(TMP, exist_ok=True)

def make_report(path, n_data, with_gauge_row=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # 表头区
    ws.cell(3, 2, "93012345")          # DWG fallback 位置
    ws.cell(3, 6, "测试件")
    ws.cell(8, 6, "Actual Readings")   # meas_start = 6
    ws.cell(9, 6, "公称尺寸")          # hdr_row = 9
    ws.cell(9, 9, "P合格")             # result_col = 9（表头区 15 行内）
    ws.cell(10, 6, 1); ws.cell(10, 7, 2); ws.cell(10, 8, 3)  # seq_row = 10
    ws.cell(9, 3, "Total Qty"); ws.cell(9, 4, n_data)
    ws.cell(9, 5, "Sample Qty"); ws.cell(8, 5, n_data)
    # 数据区：11..(10+n_data) 行
    for i in range(n_data):
        r = 11 + i
        ws.cell(r, 1, f"ITEM_{i+1}")
        ws.cell(r, 2, "间距")
        ws.cell(r, 6, 10 + i * 0.1)
        ws.cell(r, 7, 11 + i * 0.1)
        ws.cell(r, 8, 12 + i * 0.1)
        ws.cell(r, 9, 10.5 + i * 0.1)
    # 检具代码行（末尾）
    if with_gauge_row:
        r = 11 + n_data
        ws.cell(r, 1, "检具代码")
        ws.cell(r, 2, "12345")
        ws.cell(r, 6, "GAUGE-001")
    wb.save(path)

def check(name, path, expect_last, expect_rows):
    info = detect_layout(path)
    sl = info.sheets[0]
    print(f"[{name}] data_first_row={sl.data_first_row} data_last_row={sl.data_last_row} "
          f"cols={sl.sample_cols} result_col={sl.result_col}")
    ok = (sl.data_last_row == expect_last)
    cv6 = sl.col_values.get(sl.sample_cols[0], [])
    ok2 = (len(cv6) == expect_rows)
    # 末值正确性
    ok3 = (len(cv6) == 0 or abs(cv6[-1] - (10 + (expect_rows - 1) * 0.1)) < 1e-9)
    print(f"   col_values[col0] 行数={len(cv6)} 末值={cv6[-1] if cv6 else None}")
    print(f"   => data_last_row {'OK' if ok else 'FAIL'}, 行数 {'OK' if ok2 else 'FAIL'}, 末值 {'OK' if ok3 else 'FAIL'}")
    return ok and ok2 and ok3

def main():
    results = []
    # 1) 500 行 + 检具代码行 -> 终点 510（检具代码行上一行）
    p1 = os.path.join(TMP, "93012345_测试件_1件_2026.1.1-0.xlsx")
    make_report(p1, 500, with_gauge_row=True)
    results.append(("500行+检具代码", check("500行+检具代码", p1, 510, 500)))
    # 2) 无检具代码行、450 行 -> 回退最后一个有内容行 = 460（不再 400 截断）
    p2 = os.path.join(TMP, "93012346_测试件2_1件_2026.1.2-0.xlsx")
    make_report(p2, 450, with_gauge_row=False)
    results.append(("450行无检具代码", check("450行无检具代码", p2, 460, 450)))
    # 3) 小表 30 行 + 检具代码 -> 40
    p3 = os.path.join(TMP, "93012347_测试件3_1件_2026.1.3-0.xlsx")
    make_report(p3, 30, with_gauge_row=True)
    results.append(("30行+检具代码", check("30行+检具代码", p3, 40, 30)))

    allok = all(x[1] for x in results)
    print("\n==== 汇总 ====")
    for n, ok in results:
        print(f"  {n}: {'PASS' if ok else 'FAIL'}")
    print("TOTAL:", "PASS" if allok else "FAIL")
    return 0 if allok else 1

if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        sys.exit(1)
