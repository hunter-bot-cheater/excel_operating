# -*- coding: utf-8 -*-
"""多来源拼接扣减测试：单份基准文件不够时，把当前文件用满，再从其它文件补足。

场景：A 剩 5 件、B 剩 10 件，任务需 6 件
  -> A 用满 5 件（改名 -0），B 取 1 件（改名 -1），合成一份 6 件新报告。
"""
import os
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)
import openpyxl  # noqa: E402
from e2e_test import build_report, build_bom, BH, cur_suffix  # noqa: E402
from extract_core.api import run_batch  # noqa: E402

TMP = os.path.join(HERE, f"sandbox_multi_{int(time.time()*1000)}")
os.makedirs(TMP, exist_ok=True)
fails = []


def section(name):
    print("=" * 66)
    print(name)


def case_split():
    section("用例A：A 剩5 + B 剩10，任务需 6 -> A 用满 + B 取1")
    d = os.path.join(TMP, "A")
    os.makedirs(d, exist_ok=True)
    fA = build_report(d, used=5, total=10)    # 剩 5
    fB = build_report(d, used=None, total=10)  # 全新，剩 10
    bom = build_bom(d, [(BH, 6)])
    s = run_batch(bom, [fA, fB], out_dir=d)
    print("ok/fail:", s["ok_count"], s["fail_count"], "| err:", s["fail_count"])
    assert s["ok_count"] == 1, s["tasks"]
    task = s["tasks"][0]
    segs = task["segments"]
    print("  segments:", [(os.path.basename(x["path"]), x["count"],
                           os.path.basename(x["source_after"])) for x in segs])
    # A 用满 -> -0；B 取 1 -> -1
    assert [x["count"] for x in segs] == [5, 1], [x["count"] for x in segs]
    assert segs[0]["source_after"].endswith("-0.xlsx"), segs[0]["source_after"]
    assert segs[1]["source_after"].endswith("-1.xlsx"), segs[1]["source_after"]
    # 输出 6 件：前 5 列来自 A 的第 6..10 件，第 6 列来自 B 的第 1 件
    out = task["output"]
    wb = openpyxl.load_workbook(out, data_only=True)
    assert len(wb.sheetnames) == 2, wb.sheetnames   # ceil(6/5)=2
    ws1 = wb.worksheets[0]
    got1 = [ws1.cell(row=12, column=6 + j).value for j in range(5)]
    expect1 = [6.12, 7.12, 8.12, 9.12, 10.12]       # A 的第 6..10 件
    ws2 = wb.worksheets[1]
    got2 = ws2.cell(row=12, column=6).value
    seq2 = [ws2.cell(row=11, column=6 + j).value for j in range(5)]
    wb.close()
    assert [round(float(v), 2) for v in got1] == expect1, f"sheet1 应为 A 的 6..10：{got1}"
    assert round(float(got2), 2) == 1.12, f"sheet2 首列应为 B 的第 1 件：{got2}"
    assert seq2 == [1, None, None, None, None], f"sheet2 首列序号应为源报告B第1件=1：{seq2}"
    print(f"  [OK] sheet1=A的6..10 {got1}；sheet2 首列=B的第1件 {got2}；序号 {seq2}")


def case_insufficient():
    section("用例B：A 剩5 + B 剩10 = 15 < 需 16 -> NO_REMAINING")
    d = os.path.join(TMP, "B")
    os.makedirs(d, exist_ok=True)
    fA = build_report(d, used=5, total=10)   # 剩 5
    fB = build_report(d, used=None, total=10)  # 剩 10
    bom = build_bom(d, [(BH, 16)])
    s = run_batch(bom, [fA, fB], out_dir=d)
    assert s["fail_count"] == 1 and s["tasks"][0]["error_code"] == "NO_REMAINING", s["tasks"]
    # 失败应不改名
    assert cur_suffix(fA) == 5 and cur_suffix(fB) is None, (cur_suffix(fA), cur_suffix(fB))
    print("  [OK] 合计不足 -> NO_REMAINING，来源未改动")


def case_three_files():
    section("用例C：三分文件 剩3 + 剩4 + 剩20，任务需 8 -> 3 用满 + 4 用满 + 从20取1")
    d = os.path.join(TMP, "C")
    os.makedirs(d, exist_ok=True)
    f1 = build_report(d, used=7, total=10)    # 剩 3
    f2 = build_report(d, used=6, total=10)    # 剩 4
    f3 = build_report(d, used=None, total=20)  # 剩 20
    bom = build_bom(d, [(BH, 8)])
    s = run_batch(bom, [f1, f2, f3], out_dir=d)
    assert s["ok_count"] == 1, s["tasks"]
    segs = s["tasks"][0]["segments"]
    assert [x["count"] for x in segs] == [3, 4, 1], [x["count"] for x in segs]
    assert segs[0]["source_after"].endswith("-0.xlsx"), segs[0]["source_after"]
    assert segs[1]["source_after"].endswith("-0.xlsx"), segs[1]["source_after"]
    assert segs[2]["source_after"].endswith("-1.xlsx"), segs[2]["source_after"]
    print("  segments:", [(os.path.basename(x["path"]), x["count"],
                           os.path.basename(x["source_after"])) for x in segs])
    print("  [OK] 3 用满 + 4 用满 + 20 取 1；f1/f2 -> -0, f3 -> -1")


def main():
    for fn in (case_split, case_insufficient, case_three_files):
        try:
            fn()
        except Exception as e:  # noqa: BLE001
            fails.append(f"{fn.__name__}: {e}")
            print("  [FAIL]", e)
    print("=" * 66)
    if fails:
        print("❌ 失败:")
        for f in fails:
            print("  -", f)
        raise SystemExit(1)
    print("✅ 多来源拼接扣减全部通过")


if __name__ == "__main__":
    main()
